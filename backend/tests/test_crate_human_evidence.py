"""Existing human evidence survives a crate export without claiming source verification."""

import csv
import io
import json
import zipfile
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook

from app.api.v1.export import write_crate
from app.cleanup import delete_user
from app.db.models import (
    Article,
    ArticlePrediction,
    ExportJob,
    HumanAnnotation,
    InferenceJob,
    PredictionStatus,
    Project,
    ProjectMember,
    ProjectMemberRole,
    Search,
    UserRole,
)
from app.searches import ScopedProject
from tests.conftest import _make_user

FIELDS = [{"key": "ai", "label": "AI", "type": "boolean"}]


@pytest.fixture
async def corpus(session):
    project = Project(title="Evidence export", fields=FIELDS)
    session.add(project)
    await session.flush()
    article = Article(project_id=project.id, pmid="42", title="Mice", abstract="Current text")
    session.add(article)
    await session.commit()
    return project, article


async def _export(session, project, search=None):
    buf = io.BytesIO()
    scoped = ScopedProject(project, search) if search else project
    await write_crate(session, scoped, "exporter@example.org", buf, search=search)
    with zipfile.ZipFile(buf) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


async def test_crate_preserves_all_human_quotes_ids_times_and_existing_tables(
    session, corpus, researcher_user
):
    project, article = corpus
    when = datetime(2026, 1, 2, 3, 4, 5, 123456, tzinfo=UTC)
    quotes = [
        '"Souris" — hébergées à deux 🐁\nN = 12',
        "  exact spacing  ",
        "duplicate",
        "duplicate",
    ]
    annotations = [
        HumanAnnotation(
            article_id=article.id,
            user_id=researcher_user.id,
            decision_json={"ai": answer, **evidence},
            notes=f"Review {i}",
            created_at=when + timedelta(seconds=i),
        )
        for i, (answer, evidence) in enumerate(
            [(True, {"evidence_quotes": quotes}), (False, {"evidence_quotes": []}), (True, {})]
        )
    ]
    session.add_all(annotations)
    await session.commit()

    files = await _export(session, project)
    record = json.loads(files["articles/pmid-42/metadata.json"])
    human = record["human"]
    assert [h["annotation_id"] for h in human] == [str(h.id) for h in annotations]
    assert [h["created_at"] for h in human] == [h.created_at.isoformat() for h in annotations]
    assert [h["evidence_quotes"] for h in human] == [quotes, [], None]
    assert [h["answers"] for h in human] == [{"ai": True}, {"ai": False}, {"ai": True}]
    assert [h["notes"] for h in human] == ["Review 0", "Review 1", "Review 2"]
    assert human[0]["reviewer_ref"] and len({h["reviewer_ref"] for h in human}) == 1
    assert researcher_user.email not in json.dumps(human)
    assert str(researcher_user.id) not in json.dumps(human)
    assert set(human[0]) == {
        "annotation_id",
        "created_at",
        "reviewer_ref",
        "answers",
        "evidence_quotes",
        "notes",
    }
    assert record["ai"] is None and record["final_status"] == "ANNOTATED"
    rows = list(csv.DictReader(io.StringIO(files["extraction.csv"].decode("utf-8-sig"))))
    assert rows[0]["ai_human"] == "True" and rows[0]["status"] == "ANNOTATED"
    assert "evidence_quotes" not in rows[0]  # existing flat CSV columns stay compatible
    workbook = load_workbook(io.BytesIO(files["results.xlsx"]))
    header, row = list(workbook["Résultats"].values)
    assert row[header.index("Annotations humaines")] == 3
    assert row[header.index("Relu par")] == researcher_user.email
    assert row[header.index("Notes humaines")] == "Review 2"
    assert b"source" in files["README.md"]


async def test_reviewer_reference_survives_account_deletion_and_is_project_scoped(
    session, corpus, researcher_user
):
    project, article = corpus
    annotation = HumanAnnotation(article_id=article.id, user_id=researcher_user.id)
    other = Project(title="Another project", fields=FIELDS)
    session.add_all([annotation, other])
    await session.flush()
    other_article = Article(project_id=other.id, pmid="42")
    session.add(other_article)
    await session.flush()
    session.add(HumanAnnotation(article_id=other_article.id, user_id=researcher_user.id))
    await session.commit()

    async def human(target):
        files = await _export(session, target)
        return json.loads(files["articles/pmid-42/metadata.json"])["human"][0]

    before = await human(project)
    assert before["reviewer_ref"] != (await human(other))["reviewer_ref"]
    await delete_user(session, researcher_user)
    await session.commit()
    after = await human(project)
    assert after == before
    assert annotation.annotator_hash not in json.dumps(after)
    # A legacy annotation with no remaining actor cannot receive invented attribution.
    annotation.annotator_hash = None
    await session.commit()
    assert (await human(project))["reviewer_ref"] is None


async def test_blind_crate_preserves_only_search_human_evidence_and_redacts_model_outputs(
    session, corpus, researcher_user
):
    project, article = corpus
    blind = Search(project_id=project.id, name="Blind", kind="blind", lot_all=True, fields=FIELDS)
    session.add(blind)
    await session.flush()
    annotation = HumanAnnotation(
        article_id=article.id,
        search_id=blind.id,
        user_id=researcher_user.id,
        blind=True,
        decision_json={"ai": False, "evidence_quotes": ["Human blind quotation"]},
    )
    session.add_all(
        [
            annotation,
            HumanAnnotation(
                article_id=article.id,
                decision_json={"ai": True, "evidence_quotes": ["OTHER_SEARCH_QUOTE"]},
            ),
            ArticlePrediction(
                article_id=article.id,
                search_id=blind.id,
                model_used="MODEL_SECRET",
                decision_json={"ai": True, "reasoning_summary": "REASONING_SECRET"},
                evidence_quotes=["MODEL_QUOTE_SECRET"],
                status=PredictionStatus.AUTO_ACCEPTED,
            ),
            InferenceJob(
                project_id=project.id,
                search_id=blind.id,
                model_used="MODEL_SECRET",
                prefilter_model="PREFILTER_SECRET",
                result_json={"detail": "JOB_RESULT_SECRET"},
            ),
        ]
    )
    await session.commit()
    files = await _export(session, project, blind)
    record = json.loads(files["articles/pmid-42/metadata.json"])
    assert record["ai"] is None
    assert len(record["human"]) == 1
    assert record["human"][0]["annotation_id"] == str(annotation.id)
    assert record["human"][0]["evidence_quotes"] == ["Human blind quotation"]
    assert record["final_status"] == "ANNOTATED"
    assert json.loads(files["project.json"])["analysis"] is None
    for name, data in files.items():
        if name.endswith(".xlsx"):
            workbook = load_workbook(io.BytesIO(data))
            data = str([list(sheet.values) for sheet in workbook]).encode()
        for hidden in (
            b"OTHER_SEARCH_QUOTE",
            b"MODEL_SECRET",
            b"REASONING_SECRET",
            b"MODEL_QUOTE_SECRET",
            b"PREFILTER_SECRET",
            b"JOB_RESULT_SECRET",
        ):
            assert hidden not in data, name


async def test_stored_blind_crate_with_two_reviewers_requires_lead_to_download(
    client, session, corpus, researcher_user, do_login, monkeypatch, tmp_path
):
    project, article = corpus
    annotator = await _make_user(
        session, "annotator@example.org", UserRole.RESEARCHER, "password-123"
    )
    session.add_all(
        [
            ProjectMember(
                project_id=project.id,
                user_id=researcher_user.id,
                role_in_project=ProjectMemberRole.LEAD,
            ),
            ProjectMember(project_id=project.id, user_id=annotator.id),
        ]
    )
    blind = Search(project_id=project.id, name="Blind", kind="blind", lot_all=True, fields=FIELDS)
    session.add(blind)
    await session.flush()
    for reviewer, quote in (
        (researcher_user, "Lead private quote"),
        (annotator, "Annotator private quote"),
    ):
        session.add(
            HumanAnnotation(
                article_id=article.id,
                search_id=blind.id,
                user_id=reviewer.id,
                blind=True,
                decision_json={"ai": True, "evidence_quotes": [quote]},
            )
        )
    await session.commit()
    # An existing archive remains protected even when the caller supplies another search.
    buf = io.BytesIO()
    await write_crate(session, ScopedProject(project, blind), "", buf, search=blind)
    (tmp_path / "stored.zip").write_bytes(buf.getvalue())
    job = ExportJob(
        project_id=project.id,
        search_id=blind.id,
        user_id=researcher_user.id,
        status="COMPLETED",
        filename="stored.zip",
    )
    session.add(job)
    await session.commit()
    monkeypatch.setenv("EXPORT_DIR", str(tmp_path))
    url = f"/api/v1/projects/{project.id}/export/crate/{job.id}/download"
    status_url = f"/api/v1/projects/{project.id}/export/crate"
    await do_login("annotator@example.org", "password-123", consent=True)
    assert (await client.get(status_url, params={"search": str(blind.id)})).status_code == 403
    assert (await client.get(url)).status_code == 403
    assert (await client.get(url, params={"search": "ignored"})).status_code == 403
    await do_login("researcher@lab.fr", "researcher-pass-123", consent=True)
    status = await client.get(status_url, params={"search": str(blind.id)})
    assert status.status_code == 200 and status.json()["download"] is True
    response = await client.get(url)
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        human = json.loads(archive.read("articles/pmid-42/metadata.json"))["human"]
    assert {h["evidence_quotes"][0] for h in human} == {
        "Lead private quote",
        "Annotator private quote",
    }
